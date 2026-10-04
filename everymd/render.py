"""Make inputs page-able: Office files and web pages become PDFs; PowerPoint notes are read natively."""

from __future__ import annotations

import shutil
import subprocess
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

from . import config

MAC_SOFFICE = Path("/Applications/LibreOffice.app/Contents/MacOS/soffice")


class HTTPResponseError(RuntimeError):
    """The requested web document returned an unsuccessful HTTP status."""


def find_soffice() -> str | None:
    """LibreOffice on PATH (Docker, Linux) or in its usual macOS app location."""
    for name in ("soffice", "libreoffice"):
        found = shutil.which(name)
        if found:
            return found
    return str(MAC_SOFFICE) if MAC_SOFFICE.exists() else None


def office_to_pdf(path: Path, workdir: Path) -> Path:
    """Render an Office file to PDF with LibreOffice. Raises if LibreOffice is missing or fails."""
    soffice = find_soffice()
    if soffice is None:
        raise FileNotFoundError("LibreOffice (soffice) is not installed")
    # A private profile avoids clashes with a running LibreOffice and stale lock files.
    profile = (workdir / "lo-profile").resolve().as_uri()
    command = [
        soffice, f"-env:UserInstallation={profile}", "--headless",
        "--convert-to", "pdf", "--outdir", str(workdir), str(path),
    ]
    done = subprocess.run(command, capture_output=True, text=True, timeout=config.SOFFICE_TIMEOUT_S)
    pdf = workdir / f"{path.stem}.pdf"
    if done.returncode != 0 or not pdf.exists():
        detail = (done.stderr or done.stdout).strip()[:300]
        raise RuntimeError(f"LibreOffice could not render {path.name}: {detail}")
    return pdf


def web_to_pdf(target: str, out_pdf: Path) -> str:
    """Print a URL or local HTML file to an A4 PDF. Returns the browser that was used.

    Playwright's sync API refuses to run inside an active asyncio loop (Jupyter has one),
    so rendering always happens on a worker thread.
    """
    with ThreadPoolExecutor(max_workers=1) as pool:
        return pool.submit(_print_page, target, out_pdf).result()


def _print_page(target: str, out_pdf: Path) -> str:
    from playwright.sync_api import sync_playwright

    with sync_playwright() as playwright:
        try:
            browser, name = playwright.chromium.launch(channel="chrome"), "chrome"
        except Exception:  # Chrome not installed (e.g. in Docker): use Playwright's Chromium
            browser, name = playwright.chromium.launch(), "chromium"
        try:
            page = browser.new_page()
            response = page.goto(target, wait_until="load", timeout=config.BROWSER_TIMEOUT_MS)
            if response is not None and response.status >= 400:
                raise HTTPResponseError(f"HTTP {response.status} while loading {target}")
            try:  # let late content settle; pages that keep polling never go quiet, so don't wait forever
                page.wait_for_load_state("networkidle", timeout=config.BROWSER_SETTLE_MS)
            except Exception:
                pass
            # tagged + outline: the PDF keeps the page's headings as bookmarks, so heading levels survive
            page.pdf(path=str(out_pdf), format="A4", print_background=True, tagged=True, outline=True)
        finally:
            browser.close()
    return name


def pptx_notes(path: Path) -> dict[int, str]:
    """Speaker notes per slide number. Rendering to PDF drops them, so read them natively."""
    from pptx import Presentation

    notes: dict[int, str] = {}
    for number, slide in enumerate(Presentation(str(path)).slides, start=1):
        if not slide.has_notes_slide:
            continue
        frame = slide.notes_slide.notes_text_frame
        text = frame.text.strip() if frame is not None else ""
        if text:
            notes[number] = text
    return notes
